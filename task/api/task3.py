from flask import Blueprint, request, jsonify, send_file
from common import call_openai_api, find_best_match
from io import BytesIO
from docx import Document
from datetime import datetime

task3 = Blueprint('task3', __name__)

# Task3 - 기획서 생성
@task3.route('/generate-marketing-plan', methods=['POST'])
def generate_marketing_plan():
    try:
        input_data = request.json
        goal = input_data.get('goal', '').strip()
        strategy = input_data.get('strategy', '').strip()
        target_audience = input_data.get('targetAudience', '').strip()
        budget = input_data.get('budget', '').strip()

        if not all([goal, strategy, target_audience, budget]):
            return jsonify({'error': '모든 필수 입력 항목(goal, strategy, targetAudience, budget)을 제공해야 합니다.'}), 400

        # OpenAI 프롬프트 생성
        prompt = (
            f"당신은 기초 화장품 브랜드 홍보 기획안 작성 전문가입니다. "
            f"아래 정보를 바탕으로 체계적이고 구체적인 기획안을 작성해주세요.\n\n"
            f"1. 목표: {goal}\n"
            f"2. 타겟층: {target_audience}\n"
            f"3. 전략: {strategy}\n"
            f"4. 예산: {budget}\n\n"
            f"항목마다 번호를 붙이고, 자세히 설명하며, 공백 제외 700자 이상으로 작성하세요."
        )

        # OpenAI API 호출
        messages = [{"role": "user", "content": prompt}]
        marketing_plan = call_openai_api(messages, max_tokens=1200)

        # 응답 반환
        return jsonify({'marketingPlan': marketing_plan})

    except Exception as e:
        print(f"Error in /generate-marketing-plan endpoint: {e}")
        return jsonify({'error': f"서버 오류 발생: {str(e)}"}), 500


# Task3 - 기획서 다운로드 (워드 파일)
@task3.route('/download-marketing-plan', methods=['POST'])
def download_marketing_plan():
    try:
        input_data = request.json
        goal = input_data.get('goal', '').strip()
        strategy = input_data.get('strategy', '').strip()
        target_audience = input_data.get('targetAudience', '').strip()
        budget = input_data.get('budget', '').strip()
        execution = input_data.get('execution', '').strip()
        conclusion = input_data.get('conclusion', '').strip()

        if not all([goal, strategy, target_audience, budget]):
            return jsonify({'error': '모든 필수 입력 항목(goal, strategy, targetAudience, budget)을 제공해야 합니다.'}), 400

        # 워드 문서 생성
        doc = Document()
        doc.add_heading('이벤트 기획서', level=1)

        # 작성 메타데이터
        doc.add_paragraph(f"작성일자: {datetime.now().strftime('%Y년 %m월 %d일')}")
        doc.add_paragraph("작성자: 마케팅 팀")

        # 이벤트 개요
        doc.add_heading('1. 이벤트 개요', level=2)
        table1 = doc.add_table(rows=3, cols=2)
        table1.style = 'Table Grid'
        table1.cell(0, 0).text = "목적"
        table1.cell(0, 1).text = goal
        table1.cell(1, 0).text = "이벤트 방식"
        table1.cell(1, 1).text = strategy
        table1.cell(2, 0).text = "예상 결과"
        table1.cell(2, 1).text = "브랜드 인지도 상승 및 신규 고객 확보"

        # 진행 계획
        doc.add_heading('2. 진행 계획', level=2)
        doc.add_paragraph("타겟층: " + target_audience)
        doc.add_paragraph("예산: " + budget)
        doc.add_paragraph("실행 방법: " + execution)

        # 결론
        doc.add_heading('3. 결론', level=2)
        doc.add_paragraph(conclusion)

        # 워드 파일 저장
        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return send_file(file_stream, as_attachment=True, download_name="marketing_plan.docx")

    except Exception as e:
        print(f"Error in /download-marketing-plan endpoint: {e}")
        return jsonify({'error': f"서버 오류 발생: {str(e)}"}), 500
